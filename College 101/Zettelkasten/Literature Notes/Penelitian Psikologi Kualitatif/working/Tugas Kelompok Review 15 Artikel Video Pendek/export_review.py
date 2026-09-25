from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

here = Path(__file__).resolve().parent
source = here.parent.parent / 'Tugas Kelompok Review 15 Artikel Video Pendek.md'
out = here / 'Tugas Kelompok Review 15 Artikel Video Pendek.docx'
lines = source.read_text().splitlines()
rows = []
for line in lines:
    if line.startswith('| ') and (line[2:4].strip().isdigit() or line.startswith('| No ')):
        rows.append([x.strip() for x in line.split('|')[1:-1]])
assert len(rows) == 16 and all(len(row) == 8 for row in rows)

def plain(s):
    s = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', s)
    return s.replace('*', '').replace('`', '').replace('<br>', '\n')

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(14)
sec.page_height = Inches(8.5)
sec.top_margin = Inches(.42)
sec.bottom_margin = Inches(.42)
sec.left_margin = Inches(.38)
sec.right_margin = Inches(.38)
sec.header_distance = Inches(.2)
sec.footer_distance = Inches(.2)

style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(7.5)
style.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = p.add_run('Review Multiple Article')
run.bold = True
run.font.size = Pt(12)
p = doc.add_paragraph('Topik: Pengalaman mahasiswa dalam melanjutkan atau menghentikan scrolling video pendek pada waktu jeda belajar')
p.paragraph_format.space_after = Pt(6)

widths = [0.36, 2.36, 0.92, 1.30, 1.55, 1.55, 2.42, 2.78]
table = doc.add_table(rows=1, cols=8)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
table.autofit = False
for i,w in enumerate(widths):
    table.columns[i].width = Inches(w)
    table.rows[0].cells[i].width = Inches(w)
for i,text in enumerate(rows[0]):
    cell = table.rows[0].cells[i]
    cell.text = text
    for para in cell.paragraphs:
        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in para.runs:run.bold=True
trPr = table.rows[0]._tr.get_or_add_trPr()
repeat = OxmlElement('w:tblHeader')
repeat.set(qn('w:val'), 'true')
trPr.append(repeat)
for row in rows[1:]:
    cells = table.add_row().cells
    for i,content in enumerate(row):
        cells[i].width = Inches(widths[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[i].text = plain(content)
        for para in cells[i].paragraphs:
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.02

start = lines.index('## Daftar pustaka')+1
end = lines.index('## DOI untuk akses artikel')
doc.add_page_break()
p=doc.add_paragraph('Daftar pustaka')
p.style='Heading 1'
for line in lines[start:end]:
    if line.strip():
        p=doc.add_paragraph(plain(line))
        p.paragraph_format.left_indent=Inches(.22)
        p.paragraph_format.first_line_indent=Inches(-.22)
        p.paragraph_format.space_after=Pt(3)

doc.save(out)
print(out)
