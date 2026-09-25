from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

here = Path(__file__).resolve().parent
name = 'Tugas Kelompok Review Konsep Teoretis Scrolling Video Pendek'
source = here.parent.parent / f'{name}.md'
out = here / f'{name}.docx'
lines = source.read_text().splitlines()
rows = [[x.strip() for x in l.split('|')[1:-1]] for l in lines if l.startswith('| ') ]
assert len(rows) == 5 and all(len(r) == 6 for r in rows)

def plain(s):
    s = re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', s)
    return s.replace('*', '').replace('`', '').replace('<br>', '\n')

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(14), Inches(8.5)
for m in ('top_margin', 'bottom_margin'): setattr(sec, m, Inches(.42))
for m in ('left_margin', 'right_margin'): setattr(sec, m, Inches(.38))
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(7.5)
style.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph(); p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r = p.add_run('Review Multiple Literature'); r.bold = True; r.font.size = Pt(12)
p = doc.add_paragraph('(Untuk menunjang kebutuhan pemenuhan konsep teoretis)'); p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
p = doc.add_paragraph('Topik: Pengalaman mahasiswa dalam melanjutkan atau menghentikan scrolling video pendek pada waktu jeda belajar')
p.paragraph_format.space_after = Pt(6)

widths = [1.45, 1.45, 1.85, 2.55, 3.35, 2.59]
t = doc.add_table(rows=1, cols=6); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'; t.autofit = False
headers = ['Nama Teori', 'Sumber Referensi', 'Pencetus', 'Definisi Teori', 'Aspek-aspek Psikologis', 'Orientasi (Simpulkan secara mandiri)']
for i, w in enumerate(widths):
    t.columns[i].width = Inches(w)
    c = t.rows[0].cells[i]; c.width = Inches(w); c.text = headers[i]
    for para in c.paragraphs:
        para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        for run in para.runs: run.bold = True
trPr = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
for row in rows[1:]:
    cells = t.add_row().cells
    for i, content in enumerate(row):
        cells[i].width = Inches(widths[i]); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[i].text = plain(content)

def section(title, start_marker, end_marker=None):
    s = lines.index(start_marker) + 1
    e = lines.index(end_marker) if end_marker else len(lines)
    p = doc.add_paragraph(title); p.style = 'Heading 1'
    for l in lines[s:e]:
        if l.strip():
            yield doc.add_paragraph(plain(l))

doc.add_page_break()
for p in section('Keterkaitan antarkonsep', '## Keterkaitan antarkonsep', '## Daftar pustaka'):
    p.paragraph_format.space_after = Pt(6); p.runs[0].font.size = Pt(10)
for p in section('Daftar pustaka', '## Daftar pustaka'):
    p.paragraph_format.left_indent = Inches(.3); p.paragraph_format.first_line_indent = Inches(-.3)
    p.paragraph_format.space_after = Pt(4); p.runs[0].font.size = Pt(10)
doc.save(out); print(out)
